"""Fix backup slides 18-22: replace python-pptx body boxes with clones of the
slide-1 body shape (byte-identical props), then set backup bullet text.
usage: python3 deck/assets/fix_backup_bodies.py"""
import copy
from pathlib import Path
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor

ROOT = Path(__file__).resolve().parents[2]
P = ROOT / "deck" / "unified.pptx"
WHITE = RGBColor(0xF0, 0xF0, 0xF0)

BULLETS = [
    ["VRAM held elsewhere.", "Spark runs it live.", "Never kill sibling jobs."],
    ["Mutations need live SSO.", "aws login, retry once.", "Read-only demos continue."],
    ["Samples ship in git.", "Same scripts, smaller slice.", "Say so up front."],
    ["Prefix empty until deploy.", "Dist verified pre-show.", "Curl proves it after."],
    ["ComfyUI holds VRAM.", "Coordinate pre-show.", "Spark covers gaps."],
]

p = Presentation(str(P))
src = None
for sh in p.slides[0].shapes:
    if sh.has_text_frame and "Every player" in sh.text:
        src = sh._sp
        break
assert src is not None

max_id = max(sh.shape_id for s in p.slides for sh in s.shapes)
for idx in range(17, 22):
    s = p.slides[idx]
    mine = None
    for sh in s.shapes:
        if sh.has_text_frame and BULLETS[idx - 17][0] in sh.text:
            mine = sh
            break
    assert mine is not None, idx
    # drop my box, clone slide-1 body box
    sp = mine._sp
    sp.addprevious(copy.deepcopy(src))
    clone = sp.getprevious()
    max_id += 1
    clone.find(".//{*}cNvPr").set("id", str(max_id))
    sp.getparent().remove(sp)
    # set bullet text on the clone
    tf = None
    for sh in s.shapes:
        if sh.has_text_frame and "Every player" in sh.text:
            tf = sh.text_frame
            break
    tf.clear()
    for j, b in enumerate(BULLETS[idx - 17]):
        para = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
        para.text = b
        para.space_after = Pt(34)
        r = para.runs[0]
        r.font.size = Pt(24)
        r.font.color.rgb = WHITE
        r.font.name = "Arial"
p.save(str(P))
print("fixed")
