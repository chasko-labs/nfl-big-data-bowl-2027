"""Append 5 backup slides (18-22) to deck/unified.pptx. usage: python3 deck/assets/append_backup.py."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

NAVY = RGBColor(0x00, 0x00, 0x2A)
LAV = RGBColor(0xD7, 0xC7, 0xEE)
WHITE = RGBColor(0xF0, 0xF0, 0xF0)
GOLD = RGBColor(0xC9, 0xA2, 0x3F)

ROOT = __import__("pathlib").Path(__file__).resolve().parents[2]
P = ROOT / "deck" / "unified.pptx"

SLIDES = [
    ("Glimmer stays dark.",
     ["VRAM held elsewhere.",
      "Spark runs it live.",
      "Never kill sibling jobs."],
     "s12_glimmer.png",
     "Say: the local box is busy, so the cloud covers this one. The recorded cold-start plus contention numbers stand in."),
    ("SSO expired, full stop.",
     ["Mutations need live SSO.",
      "aws login, retry once.",
      "Read-only demos continue."],
     "s09_terminal.png",
     "Say: auth expired, I stop rather than route around it. Everything read-only keeps running."),
    ("Sweep dies, samples live.",
     ["Samples ship in git.",
      "Same scripts, smaller slice.",
      "Say so up front."],
     "s04_schema.png",
     "Say: the full sweep broke, so here is the same pipeline on the in-git sample slice."),
    ("404 is expected.",
     ["Prefix empty until deploy.",
      "Dist verified pre-show.",
      "Curl proves it after."],
     "s06_pipeline.png",
     "Say: a 404 before first deploy is the correct answer. Watch it flip after the sync."),
    ("VRAM held, show goes on.",
     ["ComfyUI holds VRAM.",
      "Coordinate pre-show.",
      "Spark covers gaps."],
     "s13_shootout.png",
     "Say: the lock only works if everyone honors it. Pre-show coordination is the real fix."),
]


def add_text(slide, left, top, width, height, text, size, color, name, bold=False):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    r = p.runs[0]
    r.font.size = Pt(size)
    r.font.color.rgb = color
    r.font.name = name
    r.font.bold = bold
    return box


p = Presentation(str(P))
blank = p.slide_layouts[6]
for i, (title, bullets, asset, notes) in enumerate(SLIDES):
    n = len(p.slides) + 1
    s = p.slides.add_slide(blank)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = NAVY
    add_text(s, 0.7, 0.35, 11.9, 1.0, title, 36, LAV, "Cinzel", True)
    box = s.shapes.add_textbox(Inches(0.7), Inches(2.1), Inches(4.2), Inches(4.6))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    for j, b in enumerate(bullets):
        para = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
        para.text = b
        para.space_after = Pt(12)
        r = para.runs[0]
        r.font.size = Pt(24)
        r.font.color.rgb = WHITE
        r.font.name = "Arial"
    s.shapes.add_picture(str(ROOT / "deck" / "assets" / asset),
                         Inches(5.43), Inches(2.3), Inches(7.3), Inches(4.11))
    add_text(s, 12.5, 6.9, 0.6, 0.33, str(n), 12, GOLD, "Calibri")
    for sh in s.notes_slide.shapes:
        if sh.has_text_frame and sh.text.strip() == "":
            sh.text = notes
            break
p.save(str(P))
print("slides now:", len(p.slides))
